from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import parse_qs, unquote, urljoin, urlparse


FOREIGN_ID_RE = re.compile(r"^([A-Z]{2})([0-9]+)([A-Z][0-9]?)?$")
MAX_INPUT_ITEMS = 500
MAX_INPUT_FILE_BYTES = 50 * 1024 * 1024

AMENDMENT_OPINION_PATTERNS = (
    "보정서",
    "의견서",
)

ADMIN_DOCUMENT_NAME_PATTERN = re.compile(
    r"(출원서|명세서|보정서|의견서|통지서|결정서|신고서|신청서|청구서|"
    r"제출서|답변서|진술서|번역문|증명서|도면|위임장|납부서)"
)

OPEN_DOCUMENT_CALL_RE = re.compile(
    r"\bopenDocument\s*\((?P<arguments>.*?)\)", re.IGNORECASE | re.DOTALL
)
JAVASCRIPT_STRING_RE = re.compile(r"(['\"])(.*?)\1", re.DOTALL)


def is_excluded_admin_document_title(title: str) -> bool:
    """Exclude application-form amendments while retaining specification amendments."""
    compact = re.sub(r"[\s\[\](){}·ㆍ._-]+", "", str(title or ""))
    return "출원서등보정보정서" in compact


def kipris_document_candidate_url(source: str, base_url: str = "") -> str | None:
    """Return a safe KIPRIS document-resource URL, excluding viewer/tracker HTML."""
    value = (source or "").strip()
    if not value or value.lower().startswith(("javascript:", "data:")):
        return None
    if value.startswith("blob:"):
        blob_origin = urlparse(value[5:])
        blob_host = (blob_origin.hostname or "").lower().rstrip(".")
        if blob_origin.scheme.lower() != "https":
            return None
        if not (
            blob_host == "kipris.or.kr"
            or blob_host.endswith(".kipris.or.kr")
        ):
            return None
        return value

    absolute_url = urljoin(base_url, value)
    parsed = urlparse(absolute_url)
    host = (parsed.hostname or "").lower().rstrip(".")
    if parsed.scheme.lower() != "https":
        return None
    if not (host == "kipris.or.kr" or host.endswith(".kipris.or.kr")):
        return None

    path = parsed.path.lower().rstrip("/")
    # This is the HTML popup shell. A plain GET returns markup, not the PDF.
    if path.endswith("/khome/detail/document.do"):
        return None

    marker = f"{path}?{parsed.query}".lower()
    if not any(
        token in marker
        for token in (
            ".pdf",
            "download",
            "original",
            "viewer",
            "file",
            "image",
            "document",
        )
    ):
        return None
    return absolute_url


def is_downloadable_admin_control(
    name: str,
    href: str,
    onclick: str,
    title: str,
    row_text: str,
    tag_name: str = "",
) -> bool:
    if tag_name.lower() == "a" and bool(name.strip()):
        return True
    marker = " ".join((href, onclick, title)).lower()
    if any(
        token in marker
        for token in ("document", "original", "viewer", "pdf", "file", "원문")
    ):
        return True
    if href.lower().startswith("javascript:"):
        return bool(ADMIN_DOCUMENT_NAME_PATTERN.search(name or row_text))
    return bool(ADMIN_DOCUMENT_NAME_PATTERN.search(name))


def parse_open_document_call(*values: str) -> tuple[str, str, str, str] | None:
    """KIPRIS openDocument(출원번호, 문서번호, 문서코드, 기타) 인수를 읽습니다."""
    for value in values:
        match = OPEN_DOCUMENT_CALL_RE.search(value or "")
        if not match:
            continue
        arguments = [
            item[1]
            for item in JAVASCRIPT_STRING_RE.findall(match.group("arguments"))
        ]
        if len(arguments) < 3:
            continue
        padded = (arguments + [""])[:4]
        application_number, document_number, document_code, extra = padded
        if not (
            normalize_number(application_number)
            and normalize_number(document_number)
            and normalize_number(document_code)
        ):
            continue
        return application_number, document_number, document_code, extra
    return None


@dataclass(frozen=True)
class NumberItem:
    number: str
    kind: str = "auto"  # auto, application, publication, registration

    @property
    def foreign_parts(self) -> tuple[str, str, str] | None:
        match = FOREIGN_ID_RE.fullmatch(self.number)
        if not match:
            return None
        country, digits, kind_code = match.groups()
        return country, digits, kind_code or ""

    @property
    def is_domestic(self) -> bool:
        parts = self.foreign_parts
        return parts is None or parts[0] == "KR"


def normalize_number(value: object) -> str:
    text = str(value or "").strip()
    return re.sub(r"[^0-9]", "", text)


def normalize_identifier(value: object) -> str:
    text = str(value or "").strip().upper()
    if re.match(r"^[A-Z]{2}", text):
        return re.sub(r"[^A-Z0-9]", "", text)
    return normalize_number(text)


def safe_error_text(error: object, limit: int = 1400) -> str:
    """로그와 CSV에서 인증정보를 제거하고 긴 브라우저 오류를 줄입니다."""
    message = str(error)
    message = re.sub(
        r"(?i)((?:accessKey|ServiceKey)=)[^&\s'\"]+",
        r"\1***",
        message,
    )
    message = re.sub(r"(?im)^\s*- cookie:.*$", "  - cookie: [보안상 숨김]", message)
    message = re.sub(r"(?i)(JSESSIONID=)[^;\s]+", r"\1***", message)
    message = re.sub(r"(?i)(_TRK_[A-Z_]+=)[^;\s]+", r"\1***", message)
    if "self-signed certificate in certificate chain" in message.lower():
        return (
            "사내 자체서명 인증서 때문에 별도 HTTP 요청이 거부되었습니다. "
            "Chrome 세션 내부 다운로드 방식으로 재시도합니다."
        )
    return message[:limit]


def normalize_kind(value: object) -> str:
    text = str(value or "").strip().lower()
    if text in {"출원", "출원번호", "application", "an"}:
        return "application"
    if text in {"공개", "공개번호", "publication", "opn"}:
        return "publication"
    if text in {
        "등록", "등록번호", "공고", "공고번호", "registration", "grant", "patent",
        "reg",
    }:
        return "registration"
    return "auto"


def parse_text_items(text: str) -> list[NumberItem]:
    items: list[NumberItem] = []
    seen: set[tuple[str, str]] = set()

    def add_item(raw_number: str, kind: str) -> None:
        number = normalize_identifier(raw_number)
        if len(number) < 7:
            return
        key = (number, kind)
        if key in seen:
            return
        seen.add(key)
        items.append(NumberItem(number=number, kind=kind))
        if len(items) > MAX_INPUT_ITEMS:
            raise ValueError(
                f"한 번에 처리할 수 있는 문헌은 최대 {MAX_INPUT_ITEMS}건입니다."
            )

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        prefix = re.match(
            r"^(출원|출원번호|공개|공개번호|등록|등록번호|공고|공고번호|"
            r"application|publication|registration|grant|patent|an|opn|reg)"
            r"\s*[:： ]\s*(.+)$",
            line,
            flags=re.IGNORECASE,
        )
        kind = "auto"
        if prefix:
            kind = normalize_kind(prefix.group(1))
            values = [part.strip() for part in re.split(r"[,;\t]+", prefix.group(2))]
        else:
            values = [part.strip() for part in re.split(r"[,;\t]+", line)]
            if values and normalize_kind(values[0]) != "auto":
                kind = normalize_kind(values.pop(0))
            elif values and normalize_kind(values[-1]) != "auto":
                kind = normalize_kind(values.pop())
        for value in values:
            add_item(value, kind)
    return items


def load_number_file(path: Path) -> list[NumberItem]:
    if path.stat().st_size > MAX_INPUT_FILE_BYTES:
        raise ValueError("입력 파일이 안전 제한(50MB)을 초과했습니다.")
    suffix = path.suffix.lower()
    if suffix in {".txt", ".csv"}:
        raw = path.read_text(encoding="utf-8-sig")
        if suffix == ".txt":
            return parse_text_items(raw)
        rows = list(csv.reader(raw.splitlines()))
        if not rows:
            return []
        items: list[NumberItem] = []
        for row_index, row in enumerate(rows):
            if not row:
                continue
            if row_index == 0 and any("번호" in cell for cell in row):
                continue
            items.extend(parse_text_items(",".join(row[:2])))
        return items
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX를 읽으려면 openpyxl 설치가 필요합니다.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        items: list[NumberItem] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
            values = [value for value in row[:2] if value is not None]
            if not values:
                continue
            if row_index == 0 and any("번호" in str(value) for value in values):
                continue
            items.extend(parse_text_items(",".join(str(value) for value in values)))
        workbook.close()
        return items
    raise ValueError("TXT, CSV 또는 XLSX 파일만 지원합니다.")


def safe_filename(value: str, limit: int = 120) -> str:
    value = re.sub(r"[\\/:*?\"<>|\r\n]+", "_", value).strip(" ._")
    value = value or "document"
    match = re.search(r"(\.[A-Za-z0-9]{1,8})$", value)
    suffix = match.group(1) if match else ""
    stem = value[: -len(suffix)] if suffix else value
    available = max(limit - len(suffix), 1)
    return stem[:available].rstrip(" .") + suffix


def canonical_document_label(title: str) -> str:
    """KIPRIS 문서명을 짧고 일관된 저장용 문서명으로 바꿉니다."""
    compact = re.sub(r"\s+", "", title or "")
    mappings = (
        ("최후의견제출통지서", "최후의견제출통지서"),
        ("최초의견제출통지서", "최초의견제출통지서"),
        ("의견제출통지서", "의견제출통지서"),
        ("거절결정서", "거절결정서"),
        ("등록결정서", "등록결정서"),
        ("특허출원서", "출원서"),
        ("특허원서", "출원서"),
        ("실용신안등록출원서", "출원서"),
        ("보정서", "보정서"),
        ("의견서", "의견서"),
        ("답변서", "답변서"),
        ("소명서", "소명서"),
        ("명세서", "명세서"),
    )
    for marker, label in mappings:
        if marker in compact:
            return label
    without_brackets = re.sub(r"\[[^\]]*\]", "", title or "").strip()
    return safe_filename(without_brackets or "행정서류", limit=50)


def document_pdf_filename(
    identifier: str,
    title: str,
    occurrence: int = 1,
    total_occurrences: int = 1,
) -> str:
    """출원번호_문서명[_순번].pdf 형식의 안정적인 파일명을 만듭니다."""
    normalized_identifier = normalize_identifier(identifier) or "문헌번호미상"
    label = canonical_document_label(title)
    suffix = f"_{occurrence}" if total_occurrences > 1 else ""
    return safe_filename(f"{normalized_identifier}_{label}{suffix}.pdf")


def blob_url_from_iframe_src(src: str) -> str:
    parsed = urlparse(src)
    values = parse_qs(parsed.query).get("file")
    if not values:
        raise RuntimeError("PDF 뷰어에서 원문 주소를 찾지 못했습니다.")
    return unquote(values[0]).split("#", 1)[0]
